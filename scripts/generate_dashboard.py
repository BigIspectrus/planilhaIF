#!/usr/bin/env python3
"""Gera o painel estático a partir da planilha oficial e valida os totais."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


PALETTE = [
    "#2a7d3f", "#c0392b", "#e6a817", "#6d4c9e", "#1a8c6e",
    "#3a9e52", "#d35400", "#8e44ad", "#1a7a6e", "#b7950b",
]
CATEGORIES = {
    "Serviços Terceirizados": ["terceirizad", "vigilancia", "apoio admin", "seguranca", "conservacao"],
    "Assistência Estudantil": ["paise", "assist estudantil", "auxilio"],
    "Alimentação": ["aliment", "refeit", "cestas", "generos", "pnae"],
    "Energia Elétrica": ["energia", "coelba", "eletric"],
    "Água e Saneamento": ["agua", "embasa", "saneamento", "potavel"],
    "Combustível & Veículos": ["combust", "veiculo", "frota", "licenciamento"],
    "Manutenção Predial": ["predial", "construc", "ar condicion", "dedetiz", "extintor", "recarga"],
    "Ração Animal": ["racao"],
    "Materiais": ["material", "materiais", "expedi", "laborat", "mat. limp", "mat. lab", "impressao", "copia", "grafico", "seguro", "correio", "divulgacao", "prossel"],
    "Bolsas e Diárias": ["pbiex", "pbic", "pibiex", "monitoria", "diaria", "bolsa", "ajuda de custo", "visita tecnica", "viagem", "eventos estudantis"],
    "Outros": [],
}


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(ch for ch in text if not unicodedata.combining(ch)).lower().strip()


def number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").strip()
    if text.startswith("#"):
        return 0.0
    text = re.sub(r"[^0-9,.-]", "", text)
    if not text:
        return 0.0
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def find_sheet(workbook, wanted: str):
    wanted_n = norm(wanted).replace(" ", "")
    exact = [name for name in workbook.sheetnames if norm(name).replace(" ", "") == wanted_n]
    if not exact:
        raise ValueError(f"Aba obrigatória não encontrada: {wanted}")
    return workbook[exact[0]]


def is_summary(name: str, service: str = "") -> bool:
    n, s = norm(name), norm(service)
    return (
        not n
        or "total" in n
        or "subtotal" in n
        or "total" in s
        or n in {"fornecedor", "servidor", "nome"}
        or "ministerio da educacao" in n
        or n.startswith("fonte")
    )


def parse_expenses(sheet, *, drop_embedded_rp: bool = False) -> list[dict]:
    rows: list[dict] = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        name = str(values[0] or "").strip()
        service = str(values[1] or "").strip()
        if is_summary(name, service):
            continue
        months = [number(v) for v in values[2:14]]
        other = number(values[14])
        total = number(values[15]) or sum(months) + other
        if total <= 0:
            continue
        rows.append({"name": name, "service": service, "months": months, "other": other, "total": total})

    if drop_embedded_rp:
        # A linha RP da Freedom já está incorporada na linha consolidada logo abaixo.
        def base(value: str) -> str:
            value = re.sub(r"\([^)]*(rp|restos a pagar)[^)]*\)", "", norm(value))
            return re.sub(r"\s+", " ", value).strip()

        consolidated = {(base(r["name"]), base(r["service"])) for r in rows if "rp" not in norm(r["service"]) and "restos a pagar" not in norm(r["name"])}
        rows = [
            r for r in rows
            if not (
                "restos a pagar" in norm(r["name"])
                and (base(r["name"]), base(r["service"])) in consolidated
            )
        ]
    return rows


def category(service: str, name: str = "") -> str:
    text = norm(f"{name} {service}")
    if any(key in text for key in ["aliment", "refeit", "cestas", "generos", "pnae"]):
        return "Alimentação"
    if "materiais eletric" in text:
        return "Materiais"
    for name, keys in CATEGORIES.items():
        if keys and any(key in text for key in keys):
            return name
    return "Outros"


def totals_by_category(rows: list[dict]) -> dict[str, float]:
    totals = {name: 0.0 for name in CATEGORIES}
    for row in rows:
        key = category(row["service"], row["name"])
        totals[key] += row["total"]
    return totals


def parse_budget(sheet) -> tuple[float, float, float]:
    total = custeio = assistance = 0.0
    in_assistance = False
    for values in sheet.iter_rows(values_only=True):
        label = norm(values[0])
        value = number(values[1])
        if "assistencia estudantil - 2026" in label:
            in_assistance = True
            continue
        if label == "total geral":
            total = value
        elif "total geral custeio + restos a pagar" in label:
            custeio = value
        elif in_assistance and label == "total":
            assistance = value
            in_assistance = False
    if min(total, custeio, assistance) <= 0:
        raise ValueError("Não foi possível identificar os totais do Orçamento 2026.")
    return total, custeio, assistance


def parse_planning(sheet) -> dict[str, float]:
    totals = {name: 0.0 for name in CATEGORIES}
    for values in sheet.iter_rows(min_row=3, values_only=True):
        supplier = str(values[0] or "").strip()
        service = str(values[1] or "").strip()
        if is_summary(supplier, service):
            continue
        value = number(values[3])
        if value > 0:
            totals[category(service, supplier)] += value
    return totals


def parse_diarias(sheet) -> list[dict]:
    rows: list[dict] = []
    kind = "Servidor"
    for values in sheet.iter_rows(min_row=3, values_only=True):
        name = str(values[0] or "").strip()
        if "colaboradores eventuais" in norm(name):
            kind = "Colaborador eventual"
            continue
        service = str(values[1] or "").strip()
        if is_summary(name, service):
            continue
        months = [number(v) for v in values[2:14]]
        total = number(values[15]) or sum(months)
        if total > 0:
            rows.append({"name": name, "motivo": service, "tipo": kind, "months": months, "total": total})
    return sorted(rows, key=lambda row: row["total"], reverse=True)


def parse_utilities(sheet) -> tuple[list[float], list[float]]:
    water = [0.0] * 12
    energy = [0.0] * 12
    for values in sheet.iter_rows(min_row=3, max_row=12, values_only=True):
        service = norm(values[1])
        target = water if "agua" in service else energy if "energia" in service or "eletric" in service else None
        if target is not None:
            for i, value in enumerate(values[2:14]):
                target[i] += number(value)
    return water, energy


def build_data(workbook, source_name: str, source_modified: str) -> tuple[dict, dict]:
    sheet_2026 = find_sheet(workbook, "Despesas Pagas 2026")
    rows_2026 = parse_expenses(sheet_2026, drop_embedded_rp=True)
    rows_2025 = parse_expenses(find_sheet(workbook, "Despesas 2025"))

    control_2026 = control_custeio = control_assistance = control_pnae = 0.0
    for values in sheet_2026.iter_rows(values_only=True):
        label = norm(values[0])
        if label == "total geral":
            control_2026 = number(values[15])
        elif label == "total geral custeio":
            control_custeio = number(values[15])
        elif label == "total assistencia estudantil":
            control_assistance = number(values[15])
        elif label == "total alimentacao escolar":
            control_pnae = number(values[15])
    calculated_2026 = sum(row["total"] for row in rows_2026)
    if not control_2026 or abs(calculated_2026 - control_2026) > 0.02:
        raise ValueError(f"Falha de conciliação 2026: linhas={calculated_2026:.2f}, Total Geral={control_2026:.2f}")

    orc_total, orc_custeio, orc_ae = parse_budget(find_sheet(workbook, "Orçamento 2026"))
    plan_map = parse_planning(find_sheet(workbook, "Planejamento 2026"))
    diarias = parse_diarias(find_sheet(workbook, "Diárias"))
    water_2026, energy_2026 = parse_utilities(find_sheet(workbook, "ÀguaEnergia"))

    def monthly(rows: list[dict]) -> list[float]:
        return [sum(row["months"][month] for row in rows) for month in range(12)]

    water_2025 = [0.0] * 12
    energy_2025 = [0.0] * 12
    for row in rows_2025:
        row_category = category(row["service"], row["name"])
        target = water_2025 if row_category == "Água e Saneamento" else energy_2025 if row_category == "Energia Elétrica" else None
        if target is not None:
            for i, value in enumerate(row["months"]):
                target[i] += value

    suppliers: dict[str, dict] = {}
    for row in rows_2026:
        item = suppliers.setdefault(row["name"], {"name": row["name"], "service": row["service"], "total": 0.0})
        item["total"] += row["total"]

    cat_keys = list(CATEGORIES)
    source_date = None
    if source_modified:
        try:
            source_date = datetime.fromisoformat(source_modified.replace("Z", "+00:00"))
        except ValueError:
            source_date = None
    displayed_date = source_date or datetime.now(timezone.utc)
    generated = displayed_date.astimezone(ZoneInfo("America/Bahia")).strftime("%d/%m/%Y %H:%M")
    data = {
        "planilhaNome": source_name,
        "sourceSheet": sheet_2026.title.strip(),
        "sourceModifiedTime": source_modified,
        "geradoEm": generated,
        "total2025": sum(row["total"] for row in rows_2025),
        "total2026": calculated_2026,
        "orcTotal": orc_total,
        "orcCusteio": orc_custeio,
        "orcAe": orc_ae,
        "execCusteio": control_custeio,
        "execAe": control_assistance,
        "execPnae": control_pnae,
        "month2025": monthly(rows_2025),
        "month2026": monthly(rows_2026),
        "catKeys": cat_keys,
        "catColors": PALETTE + ["#6b7280"],
        "catTotals2025": totals_by_category(rows_2025),
        "catTotals2026": totals_by_category(rows_2026),
        "planMap": plan_map,
        "suppliers": sorted(suppliers.values(), key=lambda item: item["total"], reverse=True),
        "rows2025": rows_2025,
        "rows2026": rows_2026,
        "diarias": diarias,
        "agua2025": water_2025,
        "agua2026": water_2026,
        "ene2025": energy_2025,
        "ene2026": energy_2026,
    }
    audit = {
        "sourceModifiedTime": source_modified,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceSheet": sheet_2026.title.strip(),
        "rows2026": len(rows_2026),
        "total2026": round(calculated_2026, 2),
        "controlTotal2026": round(control_2026, 2),
        "budget2026": round(orc_total, 2),
    }
    return data, audit


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--template", type=Path, default=Path("template.html"))
    parser.add_argument("--output", type=Path, default=Path("index.html"))
    parser.add_argument("--status-output", type=Path, default=Path("dashboard-source.json"))
    parser.add_argument("--source-modified", default="")
    parser.add_argument("--source-name", default="")
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, data_only=True, read_only=False)
    data, audit = build_data(workbook, args.source_name or args.workbook.name, args.source_modified)
    template = args.template.read_text(encoding="utf-8")
    if "__DATA_JSON__" not in template:
        raise ValueError("Placeholder __DATA_JSON__ ausente do template.")
    embedded_json = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")
    output = template.replace("__DATA_JSON__", embedded_json)
    output = output.replace("'__LOGO_B64__'", "''")
    args.output.write_text(output, encoding="utf-8")
    args.status_output.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    main()
